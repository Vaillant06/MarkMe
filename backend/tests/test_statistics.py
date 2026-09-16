import os
import io
import copy
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.excel.statistics import calculate_subject_statistics
from app.models.schemas import SubjectStatisticsResponse
from app.auth.session import create_session_token

WORKBOOK_PATH = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "../../V Sem B Attendance sheet.xlsx")
)

@pytest.fixture
def sample_workbook():
    assert os.path.exists(WORKBOOK_PATH), f"Fixture workbook not found at {WORKBOOK_PATH}"
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    yield wb
    wb.close()

def test_subject_statistics_uit3562(sample_workbook):
    stats = calculate_subject_statistics(
        sample_workbook,
        sheet_name="UIT3562",
        threshold=75.0,
        file_id="test_file_id",
        file_name="V Sem B Attendance sheet.xlsx"
    )

    assert isinstance(stats, SubjectStatisticsResponse)
    assert stats.sheet_name == "UIT3562"
    assert stats.subject_code == "UIT3562"
    assert stats.total_students == 71
    assert stats.total_sessions == 19
    assert stats.has_sessions is True
    assert stats.total_present == 958
    assert stats.total_absent == 391

    # Individual & average attendance calculation
    assert stats.average_attendance == 71.0
    assert 71.0 <= stats.average_attendance_raw <= 71.05

    # Below threshold calculation (threshold 75.0)
    assert stats.below_threshold_count == 41

    # Distribution checks (mutually exclusive & exhaustive)
    assert stats.distribution_90_100 == 5
    assert stats.distribution_80_89 == 16
    assert stats.distribution_75_79 == 9
    assert stats.distribution_below_75 == 41
    assert (
        stats.distribution_90_100
        + stats.distribution_80_89
        + stats.distribution_75_79
        + stats.distribution_below_75
    ) == 71

    # Overall part-to-whole percentages
    assert stats.present_percentage == 71.0
    assert stats.absent_percentage == 29.0
    assert stats.present_percentage + stats.absent_percentage == 100.0

    # Sessions trend checks
    assert len(stats.sessions) == 19
    first_session = stats.sessions[0]
    assert first_session.col_idx == 6
    assert first_session.col_letter == "F"
    assert first_session.present_count == 70
    assert first_session.absent_count == 1
    assert first_session.attendance_percentage == 98.6

def test_subject_statistics_uit3561(sample_workbook):
    stats = calculate_subject_statistics(
        sample_workbook,
        sheet_name="UIT3561",
        threshold=75.0
    )
    assert stats.sheet_name == "UIT3561"
    assert stats.total_students == 71
    assert stats.total_sessions == 5
    assert stats.total_present == 181
    assert stats.total_absent == 174
    assert stats.has_sessions is True

def test_subject_with_no_sessions(sample_workbook):
    # UIT3515 has 0 recorded attendance sessions in reference sheet
    stats = calculate_subject_statistics(
        sample_workbook,
        sheet_name="UIT3515",
        threshold=75.0
    )
    assert stats.sheet_name == "UIT3515"
    assert stats.total_students == 71
    assert stats.total_sessions == 0
    assert stats.has_sessions is False
    assert stats.average_attendance == 0.0
    assert stats.total_present == 0
    assert stats.total_absent == 0
    assert stats.below_threshold_count == 0
    assert stats.message == "No attendance sessions have been recorded for this subject."
    assert len(stats.sessions) == 0

def test_configurable_threshold(sample_workbook):
    stats_75 = calculate_subject_statistics(sample_workbook, "UIT3562", threshold=75.0)
    stats_80 = calculate_subject_statistics(sample_workbook, "UIT3562", threshold=80.0)
    stats_60 = calculate_subject_statistics(sample_workbook, "UIT3562", threshold=60.0)

    assert stats_80.below_threshold_count >= stats_75.below_threshold_count
    assert stats_75.below_threshold_count >= stats_60.below_threshold_count

def test_workbook_not_modified_during_statistics():
    with open(WORKBOOK_PATH, "rb") as f:
        original_bytes = f.read()

    wb = openpyxl.load_workbook(io.BytesIO(original_bytes), data_only=True)
    _ = calculate_subject_statistics(wb, "UIT3562")
    wb.close()

    # Re-verify original file is bitwise untouched
    with open(WORKBOOK_PATH, "rb") as f:
        after_bytes = f.read()

    assert original_bytes == after_bytes, "Excel workbook was modified during statistics calculation!"

def test_invalid_sheet_name_raises(sample_workbook):
    with pytest.raises(ValueError, match="not found in workbook"):
        calculate_subject_statistics(sample_workbook, "NON_EXISTENT_SHEET")

def test_non_subject_sheet_raises(sample_workbook):
    with pytest.raises(ValueError, match="not an active subject attendance sheet"):
        calculate_subject_statistics(sample_workbook, "SUBJECT SUMMARY SHEET")

def test_blank_and_unexpected_values():
    # Build a synthetic workbook to test blank & unexpected value counting
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TESTSUB"

    # Set headers
    ws.cell(4, 1, "Class & Section : V SEM IT B")
    ws.cell(5, 1, "Academic Year : 2026-2027")
    ws.cell(7, 1, "Sub Code & Sub Name : TEST101 - Test Subject")
    ws.cell(8, 1, "Staff Name : Staff")
    ws.cell(11, 3, "Register Number")
    ws.cell(11, 4, "Name as per 12th Mark Sheet")

    # Session headers in Row 12 (Cols 6 & 7)
    ws.cell(12, 6, "10/09/2026 1st Hour")
    ws.cell(12, 7, "11/09/2026 2nd Hour")
    # Summary formula in Col 8
    ws.cell(12, 8, "=COUNTA(F12:G12)")

    # 3 students in Rows 13, 14, 15
    students_data = [
        (13, "3122245002001", "Student One", "P", None),       # Row 13: 1 P, 1 Blank
        (14, "3122245002002", "Student Two", "A", "P"),        # Row 14: 1 A, 1 P
        (15, "3122245002003", "Student Three", "OD", "P"),     # Row 15: 1 Unexpected ('OD'), 1 P
    ]
    for r, reg, name, c6, c7 in students_data:
        ws.cell(r, 1, r - 12)
        ws.cell(r, 3, reg)
        ws.cell(r, 4, name)
        ws.cell(r, 6, c6)
        ws.cell(r, 7, c7)

    stats = calculate_subject_statistics(wb, "TESTSUB", threshold=75.0)

    assert stats.total_students == 3
    assert stats.total_sessions == 2
    assert stats.total_present == 3    # Student 1 col6, Student 2 col7, Student 3 col7
    assert stats.total_absent == 1     # Student 2 col6
    assert stats.total_unrecorded >= 2 # 1 blank + 1 unexpected ('OD')
    assert stats.unexpected_values_count == 1

    # Student 1: 1 P / 1 recorded = 100%
    # Student 2: 1 P / 2 recorded = 50%
    # Student 3: 1 P / 1 recorded ('OD' not P/A) = 100%
    # Average: (100 + 50 + 100) / 3 = 83.3%
    assert stats.average_attendance == 83.3
    assert stats.below_threshold_count == 1  # Student 2 (50% < 75%)

def test_api_statistics_endpoint_unauthorized():
    client = TestClient(app)
    res = client.get("/api/workbooks/local_test/subjects/UIT3562/statistics")
    assert res.status_code in [401, 403]

def test_api_statistics_endpoint_authorized():
    client = TestClient(app)
    # Create test session
    session_token = create_session_token({
        "id": "fac_123",
        "email": "faculty@ssn.edu.in",
        "name": "Dr. Faculty",
        "is_authorized": True
    })

    # In dev mode, 'local' file_id resolves to reference workbook
    res = client.get(
        "/api/workbooks/local/subjects/UIT3562/statistics?threshold=75.0",
        headers={"Authorization": f"Bearer {session_token}"}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["sheet_name"] == "UIT3562"
    assert data["total_students"] == 71
    assert data["total_sessions"] == 19
    assert data["average_attendance"] == 71.0
    assert len(data["sessions"]) == 19
