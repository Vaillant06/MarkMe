import shutil
import tempfile
import os
from pathlib import Path
import openpyxl
import pytest
from app.excel.updater import apply_attendance_update, DuplicateSessionError
from app.excel.parser import extract_students, extract_existing_sessions, find_summary_col

REFERENCE_FILE = str(Path(__file__).parent / "fixtures" / "sample_workbook.xlsx")

@pytest.fixture
def temp_wb_path():
    # Copy workbook to temporary file for isolated test
    temp_dir = tempfile.mkdtemp()
    temp_path = os.path.join(temp_dir, "test_attendance.xlsx")
    shutil.copyfile(REFERENCE_FILE, temp_path)
    yield temp_path
    shutil.rmtree(temp_dir, ignore_errors=True)

def test_apply_attendance_update(temp_wb_path):
    wb = openpyxl.load_workbook(temp_wb_path, data_only=False)
    
    # UIT3562 had 19 sessions (F through X), next available is Y (col 25)
    target_col, total, present, absent = apply_attendance_update(
        wb=wb,
        sheet_name="UIT3562",
        date="10/09/2026",
        period="3",
        absent_suffixes=["067", "080", "114", "129"],
        allow_overwrite=False
    )
    
    assert target_col == "Y"
    assert total == 71
    assert absent == 4
    assert present == 67
    
    wb.save(temp_wb_path)
    wb.close()
    
    # Reload and verify contents
    wb_check = openpyxl.load_workbook(temp_wb_path, data_only=False)
    ws = wb_check["UIT3562"]
    
    # Check header
    header_val = ws["Y12"].value
    assert "10/09/26" in header_val
    assert "3rd Hour" in header_val
    
    # Check student attendance
    students = extract_students(ws)
    st_map = {s.suffix: s.row for s in students}
    
    # Check absent students
    assert ws.cell(row=st_map["067"], column=25).value == "A"
    assert ws.cell(row=st_map["080"], column=25).value == "A"
    assert ws.cell(row=st_map["114"], column=25).value == "A"
    assert ws.cell(row=st_map["129"], column=25).value == "A"
    
    # Check present student
    assert ws.cell(row=st_map["068"], column=25).value == "P"
    
    # Check bottom summary formulas
    assert ws["Y91"].value == 71.0
    assert ws["Y92"].value == '=COUNTIF(Y13:Y83,"P")'
    assert ws["Y93"].value == '=COUNTIF(Y13:Y83,"A")'
    assert ws["Y94"].value == '=SUM((Y92/Y91)*100)'
    assert ws["Y95"].value == '=SUM((Y93/Y91)*100)'
    
    # Verify summary column formula remains intact
    assert ws["BL12"].value == "=COUNTA(F12:BK12)"
    assert ws["BL13"].value == '=COUNTIF(F13:BK13,"P")'
    
    wb_check.close()

def test_duplicate_session_protection(temp_wb_path):
    wb = openpyxl.load_workbook(temp_wb_path, data_only=False)
    
    # UIT3562 already has 29/06/26 1st Hour in col F
    with pytest.raises(DuplicateSessionError) as exc:
        apply_attendance_update(
            wb=wb,
            sheet_name="UIT3562",
            date="29/06/2026",
            period="1",
            absent_suffixes=["067"],
            allow_overwrite=False
        )
    assert exc.value.col_letter == "F"
    
    # Now try with allow_overwrite = True
    target_col, total, present, absent = apply_attendance_update(
        wb=wb,
        sheet_name="UIT3562",
        date="29/06/2026",
        period="1",
        absent_suffixes=["067"],
        allow_overwrite=True
    )
    assert target_col == "F"
    assert absent == 1
    assert present == 70
    
    wb.close()
