import io
import gc
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, status
from starlette.concurrency import run_in_threadpool
from app.models.schemas import WorkbookDetails, WorkbookSelectRequest, SubjectStatisticsResponse
from app.auth.dependencies import get_current_user
from app.auth.session import update_session_data
from app.drive.service import DriveService
from app.excel.parser import parse_workbook
from app.excel.statistics import calculate_subject_statistics

router = APIRouter(prefix="/api/workbooks", tags=["workbooks"])

def _parse_workbook_sync(content_bytes: bytes, file_id: str, file_name: str) -> WorkbookDetails:
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=False)
    try:
        return parse_workbook(wb, file_id=file_id, file_name=file_name)
    finally:
        wb.close()
        del wb
        gc.collect()

def _calculate_stats_sync(content_bytes: bytes, sheet_name: str, threshold: float, file_id: str, file_name: str):
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
    try:
        return calculate_subject_statistics(
            wb=wb,
            sheet_name=sheet_name,
            threshold=threshold,
            file_id=file_id,
            file_name=file_name
        )
    finally:
        wb.close()
        del wb
        gc.collect()

@router.get("/{file_id}/details", response_model=WorkbookDetails)
async def get_workbook_details(file_id: str, user: dict = Depends(get_current_user)):
    """
    Downloads and analyzes the selected attendance workbook.
    Extracts class, section, academic year, subject list, student rosters, and existing sessions.
    Automatically updates user session with the selected workbook.
    """
    drive_svc = DriveService(
        access_token=user.get("access_token"),
        refresh_token=user.get("refresh_token"),
        user_id=user.get("id") or user.get("email")
    )
    try:
        content_bytes, file_name, head_rev = await run_in_threadpool(
            drive_svc.download_workbook, file_id
        )
        
        details = await run_in_threadpool(
            _parse_workbook_sync, content_bytes, file_id, file_name
        )
        del content_bytes
        gc.collect()
        
        user_id = user.get("id") or user.get("email")
        if user_id:
            update_session_data(user_id, {
                "selected_file_id": file_id,
                "selected_file_name": file_name
            })

        return details
    except FileNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Workbook not found."
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error parsing workbook: {str(e)}"
        )

@router.post("/select")
async def select_workbook(req: WorkbookSelectRequest, user: dict = Depends(get_current_user)):
    """
    Saves the selected workbook ID to the user's backend session.
    """
    user_id = user.get("id") or user.get("email")
    if user_id:
        update_session_data(user_id, {
            "selected_file_id": req.file_id,
            "selected_file_name": req.file_name
        })
    return {"success": True, "file_id": req.file_id, "file_name": req.file_name}

@router.post("/clear")
async def clear_workbook(user: dict = Depends(get_current_user)):
    """
    Clears the selected workbook from the user's backend session.
    """
    user_id = user.get("id") or user.get("email")
    if user_id:
        update_session_data(user_id, {
            "selected_file_id": None,
            "selected_file_name": None
        })
    return {"success": True, "message": "Workbook selection cleared."}

@router.get("/{file_id}/subjects/{sheet_name}/statistics", response_model=SubjectStatisticsResponse)
async def get_subject_statistics(
    file_id: str,
    sheet_name: str,
    threshold: float = 75.0,
    user: dict = Depends(get_current_user)
):
    """
    Calculates Level 1 Subject Statistics from the workbook for the specified subject sheet.
    Reuses the existing parser and does not modify the Excel workbook.
    """
    drive_svc = DriveService(
        access_token=user.get("access_token"),
        refresh_token=user.get("refresh_token"),
        user_id=user.get("id") or user.get("email")
    )
    try:
        content_bytes, file_name, _ = await run_in_threadpool(
            drive_svc.download_workbook, file_id
        )
        stats = await run_in_threadpool(
            _calculate_stats_sync, content_bytes, sheet_name, threshold, file_id, file_name
        )
        del content_bytes
        gc.collect()
        return stats
    except FileNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Workbook not found."
        )
    except ValueError as ve:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND if "not found" in str(ve).lower() else status.HTTP_400_BAD_REQUEST,
            detail=str(ve)
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error calculating subject statistics: {str(e)}"
        )


