import io
import gc
import logging
from datetime import datetime
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, status
from starlette.concurrency import run_in_threadpool
from app.models.schemas import (
    AttendancePreviewRequest,
    AttendancePreviewResponse,
    AttendanceCommitRequest,
    AttendanceCommitResponse
)
from app.auth.dependencies import get_current_user
from app.drive.service import DriveService, DriveConcurrencyError
from app.attendance.preview import generate_attendance_preview
from app.attendance.validator import SuffixValidationError
from app.excel.updater import apply_attendance_update, DuplicateSessionError, ExcelUpdateError

from app.routes.workbooks import _EXCEL_LOCK

logger = logging.getLogger("attendance_audit")
router = APIRouter(prefix="/api/attendance", tags=["attendance"])

def _preview_sync(content_bytes: bytes, sheet_name: str, date: str, period: str, raw_input: str, entry_mode: str) -> AttendancePreviewResponse:
    with _EXCEL_LOCK:
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=False)
        del content_bytes
        try:
            if sheet_name not in wb.sheetnames:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Worksheet '{sheet_name}' not found in workbook."
                )
            ws = wb[sheet_name]
            return generate_attendance_preview(
                ws=ws,
                sheet_name=sheet_name,
                date=date,
                period=period,
                student_input=raw_input,
                entry_mode=entry_mode
            )
        finally:
            wb.close()
            del wb
            gc.collect()

def _commit_sync(content_bytes: bytes, sheet_name: str, date: str, period: str, absent_suffixes: list, allow_overwrite: bool, target_col_idx):
    with _EXCEL_LOCK:
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=False)
        del content_bytes
        try:
            col_letter, total_st, present_cnt, absent_cnt = apply_attendance_update(
                wb=wb,
                sheet_name=sheet_name,
                date=date,
                period=period,
                absent_suffixes=absent_suffixes,
                allow_overwrite=allow_overwrite,
                target_col_idx=target_col_idx
            )
            out_buf = io.BytesIO()
            wb.save(out_buf)
            res_bytes = out_buf.getvalue()
            out_buf.close()
            del out_buf
            return res_bytes, col_letter, total_st, present_cnt, absent_cnt
        finally:
            wb.close()
            del wb
            gc.collect()

@router.post("/preview", response_model=AttendancePreviewResponse)
async def preview_attendance(req: AttendancePreviewRequest, user: dict = Depends(get_current_user)):
    """
    Validates absent register suffix list and generates an exact student-by-student preview.
    Detects if the session already exists and flags duplicate warning.
    """
    drive_svc = DriveService(
        access_token=user.get("access_token"),
        refresh_token=user.get("refresh_token"),
        user_id=user.get("id") or user.get("email")
    )
    try:
        content_bytes, _, _ = await run_in_threadpool(
            drive_svc.download_workbook, req.file_id
        )
        raw_input = req.student_input if req.student_input is not None else (req.absent_input or "")
        preview_res = await run_in_threadpool(
            _preview_sync, content_bytes, req.sheet_name, req.date, req.period, raw_input, req.entry_mode
        )
        del content_bytes
        gc.collect()
        return preview_res
        
    except SuffixValidationError as e:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"message": e.message, "errors": e.errors}
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"Failed to generate attendance preview: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate attendance preview: {str(e)}"
        )

@router.post("/commit", response_model=AttendanceCommitResponse)
async def commit_attendance(req: AttendanceCommitRequest, user: dict = Depends(get_current_user)):
    """
    Applies confirmed attendance to the workbook, updates summary calculations,
    preserves formulas, formatting, and uploads the revised workbook to Google Drive.
    """
    user_email = user.get("email", "unknown")
    user_id = user.get("id") or user_email
    logger.info(f"COMMIT START | User: {user_email} | File: {req.file_id} | Sheet: {req.sheet_name} | Date: {req.date} | Period: {req.period}")

    drive_svc = DriveService(
        access_token=user.get("access_token"),
        refresh_token=user.get("refresh_token"),
        user_id=user_id
    )
    try:
        logger.info(f"Downloading workbook {req.file_id}...")
        content_bytes, file_name, current_rev = await run_in_threadpool(
            drive_svc.download_workbook, req.file_id, True
        )
        logger.info(f"Downloaded workbook {file_name} ({len(content_bytes)} bytes). Updating attendance...")

        out_bytes, col_letter, total_st, present_cnt, absent_cnt = await run_in_threadpool(
            _commit_sync,
            content_bytes,
            req.sheet_name,
            req.date,
            req.period,
            req.absent_suffixes,
            req.allow_overwrite,
            req.target_col_idx
        )
        del content_bytes
        gc.collect()
        
        logger.info(f"Attendance updated locally ({len(out_bytes)} bytes). Uploading revision to Google Drive...")
        # Upload revised workbook to Google Drive with concurrency check
        new_rev = await run_in_threadpool(
            drive_svc.upload_workbook_revision,
            req.file_id,
            out_bytes,
            req.head_revision_id
        )
        del out_bytes
        gc.collect()
        
        # Audit Log Entry
        logger.info(
            f"AUDIT | User: {user_email} | File: {file_name} | Sheet: {req.sheet_name} | "
            f"Date: {req.date} | Period: {req.period} | Total: {total_st} | "
            f"Present: {present_cnt} | Absent: {absent_cnt} | Column: {col_letter} | "
            f"Timestamp: {datetime.utcnow().isoformat()} | Rev: {new_rev}"
        )
        
        return AttendanceCommitResponse(
            success=True,
            message=f"Attendance saved successfully in column {col_letter}.",
            sheet_name=req.sheet_name,
            subject_code=req.sheet_name,
            date=req.date,
            period=req.period,
            total_students=total_st,
            present_count=present_cnt,
            absent_count=absent_cnt,
            target_col_letter=col_letter,
            new_revision_id=new_rev
        )
        
    except DuplicateSessionError as e:
        logger.warning(f"Duplicate session error committing attendance: {e}")
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={
                "message": str(e),
                "is_duplicate": True,
                "col_idx": e.col_idx,
                "col_letter": e.col_letter
            }
        )
    except DriveConcurrencyError as e:
        logger.warning(f"Drive concurrency error committing attendance: {e}")
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={"message": str(e), "is_concurrency_conflict": True}
        )
    except ExcelUpdateError as e:
        logger.warning(f"Excel update error committing attendance: {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(e)
        )
    except ValueError as e:
        logger.warning(f"Validation error committing attendance: {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(e)
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"Unexpected error committing attendance for file {req.file_id}: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to commit attendance: {str(e)}"
        )
