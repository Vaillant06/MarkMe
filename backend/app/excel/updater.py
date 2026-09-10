import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from typing import List, Optional, Tuple
from app.models.schemas import StudentInfo, AttendanceCommitResponse
from app.excel.parser import extract_students, find_summary_col, extract_existing_sessions, parse_header_metadata
from app.excel.session_matcher import format_session_header, is_duplicate_session

# Common styling definitions matching the reference workbook
THIN_SIDE = Side(border_style="thin", color="000000")
STANDARD_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
STANDARD_FONT = Font(name="Arial", size=10.0, bold=False)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=False)
BOTTOM_ALIGNMENT = Alignment(horizontal="center", vertical="bottom", wrap_text=False)

class ExcelUpdateError(Exception):
    pass

class DuplicateSessionError(ExcelUpdateError):
    def __init__(self, message: str, col_idx: int, col_letter: str):
        super().__init__(message)
        self.col_idx = col_idx
        self.col_letter = col_letter

def apply_attendance_update(
    wb: openpyxl.Workbook,
    sheet_name: str,
    date: str,
    period: str,
    absent_suffixes: List[str],
    allow_overwrite: bool = False,
    target_col_idx: Optional[int] = None
) -> Tuple[str, int, int, int]:
    """
    Applies attendance marking to the specified worksheet in openpyxl workbook.
    Returns:
    - target_col_letter
    - total_students
    - present_count
    - absent_count
    """
    if sheet_name not in wb.sheetnames:
        raise ExcelUpdateError(f"Worksheet '{sheet_name}' not found in workbook.")
        
    ws = wb[sheet_name]
    students = extract_students(ws)
    if not students:
        raise ExcelUpdateError(f"No student records found in sheet '{sheet_name}'.")
        
    summary_col = find_summary_col(ws)
    existing_sessions, next_available_col = extract_existing_sessions(ws, summary_col)
    
    # Check for existing duplicate session
    matched_col_idx = None
    for s in existing_sessions:
        if is_duplicate_session(s.header_raw, date, period):
            matched_col_idx = s.col_idx
            break
            
    if matched_col_idx is not None and not allow_overwrite:
        col_let = openpyxl.utils.get_column_letter(matched_col_idx)
        raise DuplicateSessionError(
            f"Attendance session for {date} (Period {period}) already exists in column {col_let}.",
            col_idx=matched_col_idx,
            col_letter=col_let
        )
        
    # Determine target column
    if allow_overwrite and target_col_idx is not None:
        target_col = target_col_idx
    elif allow_overwrite and matched_col_idx is not None:
        target_col = matched_col_idx
    else:
        target_col = next_available_col
        
    if target_col >= summary_col:
        raise ExcelUpdateError(
            f"Cannot add session: All pre-allocated attendance slots in sheet '{sheet_name}' are filled."
        )
        
    target_col_letter = openpyxl.utils.get_column_letter(target_col)
    
    # 1. Write Header in Row 12
    header_text = format_session_header(date, period)
    header_cell = ws.cell(row=12, column=target_col)
    header_cell.value = header_text
    header_cell.font = STANDARD_FONT
    header_cell.alignment = HEADER_ALIGNMENT
    header_cell.border = STANDARD_BORDER
    
    # Ensure column width is set
    ws.column_dimensions[target_col_letter].width = 13.0
    
    # 2. Write Student Attendance (Row 13 down)
    absent_set = set(absent_suffixes)
    present_count = 0
    absent_count = 0
    
    for st in students:
        r = st.row
        is_absent = st.suffix in absent_set
        cell = ws.cell(row=r, column=target_col)
        
        if is_absent:
            cell.value = "A"
            absent_count += 1
        else:
            cell.value = "P"
            present_count += 1
            
        cell.font = STANDARD_FONT
        cell.alignment = DATA_ALIGNMENT
        cell.border = STANDARD_BORDER
        
    first_row = students[0].row
    last_row = students[-1].row
    total_students = len(students)
    
    # 3. Write Bottom Summary Formulas (Rows 91 to 95)
    # Row 91: Total No. of Students
    cell_91 = ws.cell(row=91, column=target_col)
    cell_91.value = float(total_students)
    cell_91.font = STANDARD_FONT
    cell_91.alignment = BOTTOM_ALIGNMENT
    cell_91.border = STANDARD_BORDER
    
    # Row 92: Total No. of Presentees =COUNTIF(col13:col83,"P")
    cell_92 = ws.cell(row=92, column=target_col)
    cell_92.value = f'=COUNTIF({target_col_letter}{first_row}:{target_col_letter}{last_row},"P")'
    cell_92.font = STANDARD_FONT
    cell_92.alignment = BOTTOM_ALIGNMENT
    cell_92.border = STANDARD_BORDER
    
    # Row 93: Total No. of Absentees =COUNTIF(col13:col83,"A")
    cell_93 = ws.cell(row=93, column=target_col)
    cell_93.value = f'=COUNTIF({target_col_letter}{first_row}:{target_col_letter}{last_row},"A")'
    cell_93.font = STANDARD_FONT
    cell_93.alignment = BOTTOM_ALIGNMENT
    cell_93.border = STANDARD_BORDER
    
    # Row 94: Presentees Percentage =SUM((col92/col91)*100)
    cell_94 = ws.cell(row=94, column=target_col)
    cell_94.value = f'=SUM(({target_col_letter}92/{target_col_letter}91)*100)'
    cell_94.font = STANDARD_FONT
    cell_94.alignment = BOTTOM_ALIGNMENT
    cell_94.border = STANDARD_BORDER
    
    # Row 95: Absentees Percentage =SUM((col93/col91)*100)
    cell_95 = ws.cell(row=95, column=target_col)
    cell_95.value = f'=SUM(({target_col_letter}93/{target_col_letter}91)*100)'
    cell_95.font = STANDARD_FONT
    cell_95.alignment = BOTTOM_ALIGNMENT
    cell_95.border = STANDARD_BORDER
    
    return target_col_letter, total_students, present_count, absent_count
